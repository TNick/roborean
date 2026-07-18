"""Tests for the XLSX document driver."""

import io

from openpyxl import Workbook
from openpyxl.cell.cell import MergedCell
from roborean_documents_xlsx.driver import XlsxDocumentDriver
from roborean_spec import DocumentOperation


def _replace_op(name: str, value: str) -> DocumentOperation:
    """Build a replace_named_value operation for tests."""
    return DocumentOperation.model_validate(
        {
            "documentId": "estimate",
            "op": "replace_named_value",
            "name": name,
            "value": {
                "kind": "public_literal",
                "dataType": "string",
                "value": value,
            },
        }
    )


def _named_value_template_bytes() -> bytes:
    """Build a workbook with placeholder, formula, and merged cells."""
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.title = "Details"
    sheet["A1"] = "{{project_name}}"
    sheet["B1"] = '=A1&"-suffix"'
    sheet.merge_cells("A3:C3")
    sheet["A3"] = "{{region}}"

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


class TestXlsxReplaceNamedValue:
    """Named-value substitution across workbook string cells."""

    def test_replaces_placeholders_and_skips_formula_and_merged(self) -> None:
        """String placeholders change; formulas and merged tails stay safe."""
        driver = XlsxDocumentDriver()
        driver._template = _named_value_template_bytes()

        session = driver.begin_session({}, {"documentId": "estimate"})
        driver.apply_operation(session, _replace_op("project_name", "Alpha"))
        driver.apply_operation(session, _replace_op("region", "North"))
        driver.finalize(session)

        sheet = session.workbook["Details"]
        assert sheet["A1"].value == "Alpha"
        assert sheet["B1"].value == '=A1&"-suffix"'
        assert sheet["A3"].value == "North"
        assert isinstance(sheet["B3"], MergedCell)

        payload = driver.serialize(session)
        assert "{{project_name}}" not in str(payload)
        assert "{{region}}" not in str(payload)
